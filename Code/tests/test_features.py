import io
from pathlib import Path

import pytest

from hld_assistant import cli
from hld_assistant.service import Service
from hld_assistant.store import Store

DATA = Path(__file__).parent.parent / "sample_data"


@pytest.fixture()
def svc():
    s = Service(Store(":memory:"))
    s.v1 = s.ingest("p", "v1.txt", (DATA / "brake_system_hld_v1.txt").read_bytes(), "1.0")["doc_id"]
    s.v2 = s.ingest("p", "v2.txt", (DATA / "brake_system_hld_v2.txt").read_bytes(), "2.0")["doc_id"]
    s.ar = s.ingest("p", "brake.arxml", (DATA / "brake_system.arxml").read_bytes(), "1.0")["doc_id"]
    return s


def test_arxml_extraction(svc):
    ex = svc.extraction(svc.ar)
    assert set(ex.components) == {"WheelSpeedProvider", "BrakeController", "ActuatorDriver", "BrakeSystem"}
    assert ex.components["ActuatorDriver"]["kind"] == "ComplexDeviceDriverSwComponentType"
    assert {(c["from_comp"], c["to_comp"]) for c in ex.connections} == {
        ("WheelSpeedProvider", "BrakeController"), ("BrakeController", "ActuatorDriver")}
    assert {r["name"]: r["period"] for r in ex.runnables}["Run_BrakeCtrl_10ms"] == "10 ms"
    wsp = next(p for p in ex.ports if p.name == "WheelSpeedOut")
    assert (wsp.signal, wsp.dtype, wsp.interface) == ("WheelSpeed", "uint16", "IfWheelSpeed")


def test_arxml_findings_and_qa(svc):
    rules = {f["rule"] for f in svc.findings(svc.ar)}
    assert "UNCONNECTED_REQUIRE_PORT" in rules  # DiagOut is not wired
    assert svc.ask("Which component provides WheelSpeed?", [svc.ar])["citations"]


def test_invalid_arxml_rejected():
    with pytest.raises(ValueError, match="Invalid ARXML"):
        Service(Store(":memory:")).ingest("p", "x.arxml", b"<not-closed")


def test_requirements_extracted_and_scoped(svc):
    reqs = svc.requirements(svc.v1)
    timing = [r for r in reqs if r["kind"] == "timing"]
    assert len(timing) == 1 and timing[0]["limit_ms"] == 20.0


def test_timing_budget_flagged_when_exceeded(svc):
    f = [x for x in svc.findings(svc.v2) if x["rule"] == "TIMING_BUDGET_EXCEEDED"]
    assert f and "17" in f[0]["message"] and f[0]["severity"] == "info"  # unscoped requirement -> info
    assert not [x for x in svc.findings(svc.v1) if x["rule"] == "TIMING_BUDGET_EXCEEDED"]  # 5+10+2 = 17 <= 20


def test_chains_latency(svc):
    best = svc.chains(svc.v1)[0]
    assert best["path"] == ["WheelSpeedProvider", "BrakeController", "ActuatorDriver"] and best["latency_ms"] == 17


def test_impact(svc):
    r = svc.impact(svc.v2, "signal", "WheelSpeed")
    assert set(r["directly_affected"]) == {"WheelSpeedProvider", "BrakeController"}
    assert [d["component"] for d in r["downstream"]] == ["ActuatorDriver"]
    up = svc.impact(svc.v2, "component", "ActuatorDriver")["upstream"]
    assert {u["component"] for u in up} == {"BrakeController", "WheelSpeedProvider", "VehicleSpeedEstimator"}
    assert not svc.impact(svc.v2, "component", "Nope")["found"]
    with pytest.raises(ValueError):
        svc.impact(svc.v2, "bogus", "x")


def test_dependency_cycle_detected():
    s = Service(Store(":memory:"))
    txt = ("1. Components\n1.1 A (ApplicationSwComponentType)\nDepends on: B\n"
           "1.2 B (ApplicationSwComponentType)\nDepends on: A\n")
    d = s.ingest("p", "c.txt", txt.encode())["doc_id"]
    assert any(f["rule"] == "DEPENDENCY_CYCLE" for f in s.findings(d))


def test_text_diff_catches_value_change(svc):
    rows = svc.text_diff(svc.v1, svc.v2)
    vals = [v for r in rows for v in r["value_changes"]]
    assert {"old": "20 ms", "new": "15 ms"} in vals
    assert {"old": "10 ms", "new": "5 ms"} in vals


def test_test_plan_traces(svc):
    plan = svc.test_plan(svc.v2)
    assert len({t["id"] for t in plan}) == len(plan)
    assert any(t["type"] == "timing" and "15" in t["expected"] for t in plan)
    assert any(t["type"] == "boundary" and "WheelSpeedOut" in t["trace_to"] for t in plan)
    assert all(t["trace_to"] and t["section"] for t in plan)


def test_facts_answer_when_retrieval_declines(svc):
    r = svc.ask("zzz WheelSpeed", [svc.v2])
    assert r["facts"] and "WheelSpeedProvider" in r["facts"][0]["text"]
    r = svc.ask("who consumes BrakeTorqueRequest", [svc.v2])
    assert any("ActuatorDriver" in f["text"] for f in r["facts"])


def test_approval_gate(svc):
    with pytest.raises(ValueError, match="Cannot approve"):
        svc.set_status(svc.v1, "approved", "bob")
    for f in svc.findings(svc.v1):
        if f["severity"] == "error":
            svc.review_finding(svc.v1, f["id"], "confirmed", "tracked", "bob")
    svc.set_status(svc.v1, "approved", "bob")
    assert svc._doc(svc.v1)["status"] == "approved"
    with pytest.raises(ValueError):
        svc.set_status(svc.v1, "bogus")


def test_metrics(svc):
    svc.ask("period of BrakeController", [svc.v1])
    svc.ask("bluetooth pairing handshake", [svc.v1])
    f = next(x for x in svc.findings(svc.v1) if x["severity"] == "error")
    svc.review_finding(svc.v1, f["id"], "confirmed")
    m = svc.metrics("p")
    assert m["questions"] == 2 and m["answer_rate_pct"] == 50.0 and m["groundedness_pct"] == 100.0
    assert m["finding_precision_pct"] == 100.0 and m["documents"] == 3


def test_report_and_export(svc):
    html = svc.report_html(svc.v1)
    assert "Sign-off" in html and "DATATYPE_MISMATCH" in html
    import zipfile
    z = zipfile.ZipFile(io.BytesIO(svc.export_zip(svc.v1)))
    assert {"tests.csv", "requirements.csv", "review_report.html"} <= set(z.namelist())


def test_docx_ingest():
    from docx import Document

    d = Document()
    d.add_heading("Interfaces", 1)
    d.add_heading("IfX (SenderReceiver Interface)", 2)
    d.add_paragraph("- Data element Val: uint8")
    d.add_heading("Components", 1)
    d.add_heading("Comp (ApplicationSwComponentType)", 2)
    d.add_paragraph("- P-Port Out: interface IfX, signal Val, type uint8")
    buf = io.BytesIO()
    d.save(buf)
    s = Service(Store(":memory:"))
    r = s.ingest("p", "h.docx", buf.getvalue())
    assert r["components"] == 1 and r["interfaces"] == 1 and r["ports"] == 1


def test_cli(tmp_path, capsys):
    assert cli.main(["analyze", str(DATA / "brake_system_hld_v1.txt")]) == 1   # has errors
    assert cli.main(["analyze", str(DATA / "brake_system_hld_v2.txt"), "--fail-on", "error"]) == 0
    capsys.readouterr()
    cli.main(["compare", str(DATA / "brake_system_hld_v1.txt"), str(DATA / "brake_system_hld_v2.txt")])
    assert "20 ms -> 15 ms" in capsys.readouterr().out
